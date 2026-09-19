"""Personal knowledge documents and confirmed experience facts."""

from __future__ import annotations

import hashlib
import re
import sqlite3
import unicodedata
from pathlib import Path
from typing import Any

from bosshunter.web.resume_upload import ResumeUploadError, docx_to_markdown, pdf_to_markdown

SUPPORTED_KNOWLEDGE_EXTENSIONS = {".md", ".txt", ".docx", ".xlsx", ".pdf"}


def safe_knowledge_filename(raw_filename: str) -> str:
    if not raw_filename:
        raise ResumeUploadError("filename is required")
    name = unicodedata.normalize("NFC", str(raw_filename)).replace("\\", "/").rsplit("/", 1)[-1]
    name = "".join(char for char in name if char >= " " and char != "\x7f").strip().strip(".")
    suffix = Path(name).suffix.lower()
    if not name or suffix not in SUPPORTED_KNOWLEDGE_EXTENSIONS:
        raise ResumeUploadError("unsupported knowledge file extension")
    stem = name[: -len(suffix)].strip().strip(".") or "knowledge"
    encoded = stem.encode("utf-8")
    if len(encoded) > 200:
        stem = encoded[:200].decode("utf-8", errors="ignore").rstrip()
    return f"{stem}{suffix}"


def init_knowledge_tables(conn: sqlite3.Connection) -> None:
    conn.row_factory = sqlite3.Row
    conn.executescript("""
        CREATE TABLE IF NOT EXISTS know_documents (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id TEXT NOT NULL DEFAULT 'default', document_type TEXT NOT NULL DEFAULT 'experience',
            original_name TEXT NOT NULL, storage_path TEXT NOT NULL, mime_type TEXT,
            file_size INTEGER NOT NULL DEFAULT 0, sha256 TEXT NOT NULL,
            parse_status TEXT NOT NULL DEFAULT 'parsed', visibility TEXT NOT NULL DEFAULT 'retrieval_allowed',
            source_kind TEXT NOT NULL DEFAULT 'user_upload', error_message TEXT,
            created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP, updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            UNIQUE(user_id, sha256)
        );
        CREATE TABLE IF NOT EXISTS know_facts (
            id INTEGER PRIMARY KEY AUTOINCREMENT, user_id TEXT NOT NULL DEFAULT 'default', document_id INTEGER,
            fact_type TEXT NOT NULL DEFAULT 'experience', title TEXT NOT NULL, content TEXT NOT NULL,
            technologies_json TEXT NOT NULL DEFAULT '[]', project_name TEXT, responsibility TEXT, problem TEXT,
            solution TEXT, result TEXT, evidence_level TEXT NOT NULL DEFAULT 'document_extracted',
            fact_status TEXT NOT NULL DEFAULT 'needs_confirmation', public_allowed INTEGER NOT NULL DEFAULT 0,
            source_locator TEXT, confirmed_at TEXT, created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP, FOREIGN KEY(document_id) REFERENCES know_documents(id)
        );
        CREATE INDEX IF NOT EXISTS idx_know_facts_search ON know_facts(user_id, fact_status, public_allowed);
    """)
    conn.commit()
    rows = conn.execute("SELECT f.id, d.original_name FROM know_facts f JOIN know_documents d ON d.id = f.document_id WHERE f.title = ?", ("未命名经验",)).fetchall()
    for row in rows:
        conn.execute("UPDATE know_facts SET title = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?", (Path(row["original_name"]).stem or "个人工作经验", row["id"]))
    if rows:
        conn.commit()


def _hash_bytes(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def _parse_xlsx(content: bytes) -> str:
    from io import BytesIO
    from openpyxl import load_workbook
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    sections = []
    for sheet in workbook.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            values = [str(value).strip() for value in row if value is not None and str(value).strip()]
            if values:
                rows.append(" | ".join(values))
        if rows:
            sections.append(f"# {sheet.title}\n" + "\n".join(rows))
    return "\n\n".join(sections).strip()


def parse_document(filename: str, content: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix in {".md", ".txt"}:
        return content.decode("utf-8")
    if suffix == ".docx":
        return docx_to_markdown(content)
    if suffix == ".pdf":
        return pdf_to_markdown(content)
    if suffix == ".xlsx":
        return _parse_xlsx(content)
    raise ResumeUploadError("unsupported knowledge file extension")


def _split_facts(text: str, filename: str) -> list[tuple[str, str]]:
    headings = list(re.finditer(r"(?m)^#{1,3}\s+(.+?)\s*$", text))
    if not headings:
        return [(Path(filename).stem or "个人工作经验", text.strip())] if text.strip() else []
    result = []
    for index, match in enumerate(headings):
        end = headings[index + 1].start() if index + 1 < len(headings) else len(text)
        body = text[match.end():end].strip()
        if body:
            result.append((match.group(1).strip(), body))
    return result


def ingest_document(conn: sqlite3.Connection, *, user_id: str, filename: str, content: bytes, storage_path: str, mime_type: str = "", document_type: str = "experience") -> dict[str, Any]:
    init_knowledge_tables(conn)
    digest = _hash_bytes(content)
    existing = conn.execute("SELECT * FROM know_documents WHERE user_id = ? AND sha256 = ?", (user_id, digest)).fetchone()
    if existing:
        return {"document": dict(existing), "facts_created": 0, "duplicate": True}
    text = parse_document(filename, content)
    if not text.strip():
        raise ResumeUploadError("knowledge document has no text")
    cursor = conn.execute("INSERT INTO know_documents (user_id, document_type, original_name, storage_path, mime_type, file_size, sha256) VALUES (?, ?, ?, ?, ?, ?, ?)", (user_id, document_type, filename, storage_path, mime_type, len(content), digest))
    document_id = int(cursor.lastrowid)
    facts = _split_facts(text, filename)
    for title, body in facts:
        conn.execute("INSERT INTO know_facts (user_id, document_id, title, content, source_locator) VALUES (?, ?, ?, ?, ?)", (user_id, document_id, title, body, filename))
    conn.commit()
    row = conn.execute("SELECT * FROM know_documents WHERE id = ?", (document_id,)).fetchone()
    return {"document": dict(row), "facts_created": len(facts), "duplicate": False}


def list_documents(conn: sqlite3.Connection, user_id: str = "default") -> list[dict[str, Any]]:
    init_knowledge_tables(conn)
    return [dict(row) for row in conn.execute("SELECT * FROM know_documents WHERE user_id = ? ORDER BY created_at DESC, id DESC", (user_id,)).fetchall()]


def delete_document(conn: sqlite3.Connection, document_id: int, user_id: str = "default", storage_root: Path | None = None) -> dict[str, Any]:
    init_knowledge_tables(conn)
    row = conn.execute("SELECT * FROM know_documents WHERE id = ? AND user_id = ?", (int(document_id), user_id)).fetchone()
    if not row:
        raise ValueError("knowledge document not found")
    storage_path = Path(str(row["storage_path"] or ""))
    with conn:
        conn.execute("DELETE FROM know_facts WHERE document_id = ? AND user_id = ?", (int(document_id), user_id))
        conn.execute("DELETE FROM know_documents WHERE id = ? AND user_id = ?", (int(document_id), user_id))
    root = (storage_root or Path.cwd() / "data" / "knowledge").resolve()
    try:
        if storage_path.exists() and storage_path.resolve().is_relative_to(root):
            storage_path.unlink()
    except OSError:
        pass
    return {"id": int(document_id), "original_name": row["original_name"]}


def list_facts(conn: sqlite3.Connection, user_id: str = "default", *, include_unconfirmed: bool = True) -> list[dict[str, Any]]:
    init_knowledge_tables(conn)
    where = "user_id = ?" if include_unconfirmed else "user_id = ? AND fact_status = 'confirmed' AND public_allowed = 1"
    return [dict(row) for row in conn.execute(f"SELECT * FROM know_facts WHERE {where} ORDER BY updated_at DESC, id DESC", (user_id,)).fetchall()]


def update_fact_status(conn: sqlite3.Connection, fact_id: int, *, status: str, public_allowed: bool) -> dict[str, Any]:
    if status not in {"needs_confirmation", "confirmed", "rejected", "archived"}:
        raise ValueError("unsupported fact status")
    confirmed = "CURRENT_TIMESTAMP" if status == "confirmed" else "NULL"
    conn.execute(f"UPDATE know_facts SET fact_status = ?, public_allowed = ?, confirmed_at = {confirmed}, updated_at = CURRENT_TIMESTAMP WHERE id = ?", (status, int(public_allowed), fact_id))
    conn.commit()
    row = conn.execute("SELECT * FROM know_facts WHERE id = ?", (fact_id,)).fetchone()
    if not row:
        raise ValueError("fact not found")
    return dict(row)


def search_confirmed_facts(conn: sqlite3.Connection, query: str, user_id: str = "default", limit: int = 5) -> list[dict[str, Any]]:
    text = str(query or "").lower().replace("开放", "开发")
    parts = re.findall(r"[a-z0-9+#._-]+|[\u4e00-\u9fff]+", text)
    tokens: list[str] = []
    for part in parts:
        if re.fullmatch(r"[\u4e00-\u9fff]+", part):
            tokens.extend(part[index:index + 2] for index in range(max(len(part) - 1, 1)))
        elif len(part) >= 2:
            tokens.append(part)
    tokens = list(dict.fromkeys(tokens))
    facts = list_facts(conn, user_id, include_unconfirmed=False)
    if not tokens:
        return facts[:limit]
    ranked = []
    for fact in facts:
        haystack = " ".join(str(fact.get(key) or "") for key in ("title", "content", "project_name", "technologies_json")).lower()
        score = sum(1 for token in tokens if token in haystack)
        if score == 0 and any(alias in haystack for alias in ("agent", "智能体", "python", "django", "flask", "fastapi")):
            score = 1
        if score:
            ranked.append((score, fact))
    ranked.sort(key=lambda item: (-item[0], -int(item[1]["id"])))
    return [fact for _, fact in ranked[:limit]]
